from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path('/Users/ashang/code/githubFile/zc-wx-form')
OUT = ROOT / '展晨门窗系统使用说明.docx'
IMG_ROOT = Path('/var/folders/tr/8j445tz917330mggsnbzjqw00000gn/T')


def img(name):
    return IMG_ROOT / f'codex-clipboard-{name}.png'


MINI = [
    ('首页', 'd081511b-9535-4f35-b80c-dd8e85f6891e', '进入小程序后可浏览 Banner、热门定制推荐、产品分类和门店信息。点击 Banner 或热门商品可进入商品详情。'),
    ('产品列表页', '5411a9c4-01f6-489e-a35d-9c08529e0432', '按分类浏览门窗商品，也可以使用搜索框查找产品。切换分类时搜索内容会自动清空。'),
    ('产品下单页 1', 'ba9179d3-89e3-4e6d-b96f-7a09134aa2c5', '在商品详情中填写尺寸，并根据需要选择玻璃、门锁、铝材等配置。'),
    ('产品下单页 2', 'bca1ef0d-2a2e-477c-ba6c-d6de1f8c5026', '可新增多套定制方案，为每套方案填写套名、尺寸、选配和备注。'),
    ('提交订单页', '50dccf5a-c322-42d2-b745-645aaf662eb9', '确认联系人、联系电话、收货/安装地址、现场图片和备注后提交订单。提交前请核对金额和方案。'),
    ('订单列表页', '95b4f509-abe0-48b0-82bd-424692ef384a', '查看全部订单或按待复核、生产中、待提货、已完成等状态筛选。状态更新提醒会显示状态 Tag 和订单号。'),
    ('订单详情页', 'a81a9adb-d7fe-4ce8-95ac-f2e204e6e1ab', '查看订单进度时间轴、客户信息、收货/安装地址、每套门窗的尺寸、选配和金额。'),
    ('个人中心页', 'bdb5877f-e753-4650-bd0c-5241b1dd715a', '管理个人资料、地址，并按订单状态进入订单列表；也可以拨打官方电话或打开联系我们。'),
    ('编辑资料/登录页', 'b97a330f-718b-45a1-a2e2-d93ca73144b1', '首次使用或资料不完整时，填写姓名和手机号并保存。小程序会优先静默登录，只有需要补充资料时才进入此页面。'),
    ('地址列表', '181e0087-be01-4d26-98e9-3ac858bd8c23', '新增、编辑、删除收货/安装地址，并设置默认地址。下单时会优先使用默认地址。'),
    ('联系我们', 'f03a89bd-1e73-4bd5-9204-188973a2668d', '查看门店地址、营业时间和接单顾问信息，可直接拨打电话或按页面提示添加顾问。'),
]

ADMIN = [
    ('登录页', '3ab233b5-9f2a-4963-b237-b53fea3a00e9', '输入管理员账号和密码后点击“登录系统”。上线前请使用实际分配的账号，不要依赖截图中的示例内容。'),
    ('大盘数据', '3fd1a190-127e-41ae-bbd0-da8ae8a2bb04', '查看订单总数、各状态数量、近 7 天订单趋势、热门门窗占比和订单流转阶段。点击“刷新图表数据”获取最新数据。'),
    ('Banner 管理', '57bb51f5-9fde-4e0f-b770-d6b4c3ed67c4', '新增 Banner 时填写名称，可选填副标题，上传图片，按需绑定商品并设置排序和启用状态。绑定商品后，小程序点击 Banner 会直接进入商品详情。'),
    ('门窗分类', 'f95279b1-becc-4b79-822c-50c998ebe960', '维护小程序首页和产品页使用的分类名称、首页小标签、图标、排序权重和启用状态。排序数字越小越靠前。'),
    ('商品列表', '06ff7136-279c-4060-9804-ade2f1b9dd66', '搜索和筛选商品，查看基础单价、起步计费面积和选配规则数量。可独立切换上架/下架和热门推荐，也支持批量上下架与复制商品。'),
    ('新增/编辑商品', '2bc074cd-5564-46a8-a3b4-bb2be79895c0', '填写商品名称、封面图、上下架状态、所属分类、描述、默认尺寸、起步计费面积和基础单价。保存成功后商品才会出现在小程序。'),
    ('商品选配规则', '746422b5-2eae-45e5-8cda-ac73ed4bbbc9', '按选配分组维护玻璃、门锁、铝材、颜色等配置。每组选项可设置图片、加价方式、加价金额和默认勾选，保存后同步到小程序下单页。'),
    ('订单列表（可创建订单）', '88ae1316-6177-4e86-8925-9f1d858b9f81', '按订单号、客户、商品、状态和订单时间筛选。点击“创建订单”可代客户下单；列表中可查看详情、修改状态或删除订单。'),
    ('查看订单详情', '73892d62-148f-4ad5-bcfd-54792c372241', '查看订单的客户信息、方案明细、尺寸、选配、金额和备注。需要调整订单状态或追加费用时，从订单列表进入“修改状态”。'),
    ('接单员配置', '81581fcb-e4fe-4434-856f-8f36f008a0da', '维护小程序“联系我们”页面展示的接单员姓名、电话、个人微信二维码和显示状态。'),
    ('客户管理（可新增/编辑/添加地址）', 'b6f1e4bf-fe39-4e31-81b0-c27386296367', '搜索客户并查看下单数量。可创建客户、修改资料、删除客户或进入客户详情添加地址。'),
    ('查看客户详情', '409ef788-e5df-464e-85d6-12b0634e1f9b', '查看客户账号档案和已保存地址；可新增地址，便于后台代客户创建订单。'),
    ('管理员管理', 'a5b1219b-8d08-4407-ab73-4a90b623a120', '新增和维护后台管理员账号、姓名/备注、电话、所属角色和账号状态。ROOT 超级管理员账号受保护，不能随意删除。'),
    ('角色/权限', 'b5c067dc-45aa-40ec-8fb6-fa8a6edc8bdb', '创建角色并分配菜单权限。给普通管理员分配角色后，他只能访问被授权的后台功能；ROOT 角色拥有最高权限。'),
    ('操作日志', 'f12afd56-976a-487d-8e4c-c3f154a19fc1', '按事件类型、管理员账号、用户、手机号和日期范围筛选操作记录，用于追踪订单状态、商品和系统配置变更。'),
    ('全局设置', '94f2cd0f-9d9c-440c-97e9-97bdd73e1285', '维护门店名称、门店地址、客服电话、营业时间和地图经纬度。保存后会影响小程序个人中心、联系我们和地图定位。'),
]


def set_font(run, name='Hiragino Sans GB', size=11, color='263238', bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths)))
    tblW.set(qn('w:type'), 'dxa')
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:w'), str(width))
            tcW.set(qn('w:type'), 'dxa')
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_para(doc, text='', size=11, color='263238', bold=False, align=None, before=0, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_font(r, size=size, color=color, bold=bold)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    if level == 1:
        set_font(r, size=16, color='2E74B5', bold=True)
    elif level == 2:
        set_font(r, size=13, color='2E74B5', bold=True)
    else:
        set_font(r, size=12, color='1F4D78', bold=True)
    return p


def new_numbering_instance(doc):
    numbering = doc.part.numbering_part._element
    current_ids = [int(node.get(qn('w:numId'))) for node in numbering.findall(qn('w:num'))]
    num = OxmlElement('w:num')
    num.set(qn('w:numId'), str(max(current_ids, default=0) + 1))
    abstract_num_id = OxmlElement('w:abstractNumId')
    abstract_num_id.set(qn('w:val'), '7')  # The document's decimal List Number definition.
    num.append(abstract_num_id)
    level_override = OxmlElement('w:lvlOverride')
    level_override.set(qn('w:ilvl'), '0')
    start_override = OxmlElement('w:startOverride')
    start_override.set(qn('w:val'), '1')
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return int(num.get(qn('w:numId')))


def apply_numbering(paragraph, numbering_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn('w:numPr'))
    if num_pr is None:
        num_pr = OxmlElement('w:numPr')
        p_pr.append(num_pr)
    ilvl = OxmlElement('w:ilvl')
    ilvl.set(qn('w:val'), '0')
    num_id = OxmlElement('w:numId')
    num_id.set(qn('w:val'), str(numbering_id))
    num_pr.append(ilvl)
    num_pr.append(num_id)


def add_step(doc, text, numbering_id):
    p = doc.add_paragraph(style='List Number')
    apply_numbering(p, numbering_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_font(r, size=11, color='263238')
    return p


def add_note(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, 'F4F6F9')
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run('提示：' + text)
    set_font(r, size=10.5, color='5B6573')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_screenshot(doc, path, caption, mobile=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    width = Inches(2.75 if mobile else 6.25)
    run = p.add_run()
    run.add_picture(str(path), width=width)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run('图：' + caption)
    set_font(r, size=9, color='6B7280')


def add_section(doc, title, image_name, description, steps, mobile, page_break_before=False):
    heading = add_heading(doc, title, 2)
    # Break before the next title rather than after a full-page image. This avoids blank pages.
    heading.paragraph_format.page_break_before = page_break_before
    add_para(doc, description, size=11, after=5)
    numbering_id = new_numbering_instance(doc)
    for step in steps:
        add_step(doc, step, numbering_id)
    add_screenshot(doc, img(image_name), title, mobile=mobile)


def configure_styles(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    normal = doc.styles['Normal']
    normal.font.name = 'Hiragino Sans GB'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string('263238')
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ('Heading 1', 16, '2E74B5', 18, 10),
        ('Heading 2', 13, '2E74B5', 14, 7),
        ('Heading 3', 12, '1F4D78', 10, 5),
    ]:
        st = doc.styles[name]
        st.font.name = 'Hiragino Sans GB'
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.25
    for style_name in ['List Number', 'List Bullet']:
        st = doc.styles[style_name]
        st.font.name = 'Hiragino Sans GB'
        st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
        st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.375)
        st.paragraph_format.first_line_indent = Inches(-0.188)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.25
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run('展晨门窗系统使用说明')
    set_font(r, size=9, color='7B8794')
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run('展晨门窗 · 内部操作参考')
    set_font(r, size=9, color='7B8794')


def build():
    doc = Document()
    configure_styles(doc)

    add_para(doc, '展晨门窗', size=12, color='A67C52', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=48, after=14)
    add_para(doc, '小程序与管理后台\n使用说明', size=28, color='1F2937', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    add_para(doc, '面向日常业务人员的图文操作手册', size=14, color='6B7280', align=WD_ALIGN_PARAGRAPH.CENTER, after=26)
    add_note(doc, '本手册根据当前系统界面编写。小程序用于客户浏览、配置和提交订单；管理后台用于维护商品、订单、客户和系统信息。')
    add_para(doc, '版本：2026 年 8 月', size=10, color='8A94A6', align=WD_ALIGN_PARAGRAPH.CENTER, before=16, after=0)
    doc.add_page_break()

    add_heading(doc, '目录与使用说明', 1)
    add_para(doc, '本手册按实际使用顺序组织。第一次使用时，建议先阅读对应模块的“页面作用”，再按编号步骤操作。截图中的数据仅用于演示，实际账号、客户、商品和订单以系统当前数据为准。')
    toc = doc.add_table(rows=1, cols=2)
    set_table_geometry(toc, [1700, 7660])
    for c, text in zip(toc.rows[0].cells, ['模块', '包含功能']):
        set_cell_shading(c, 'E8EEF5')
        c.text = text
        for r in c.paragraphs[0].runs: set_font(r, size=10.5, color='1F4D78', bold=True)
    for label, detail in [('一、小程序端', '首页、产品列表、产品配置、提交订单、订单、个人中心、资料、地址、联系我们'), ('二、管理后台', '登录、大盘、Banner、分类、商品、选配规则、订单、接单员、客户、管理员、角色、日志、全局设置'), ('三、常见问题', '登录、数据更新、权限和操作注意事项')]:
        cells = toc.add_row().cells
        cells[0].text, cells[1].text = label, detail
        for c in cells:
            for r in c.paragraphs[0].runs: set_font(r, size=10.5, color='263238')
    add_note(doc, '按钮保存后请等待提示。若提示失败，不要重复点击，先检查网络和填写内容，再按页面提示重试。')
    doc.add_page_break()

    add_heading(doc, '一、小程序端使用说明', 1)
    add_para(doc, '小程序面向客户使用，主要流程是：浏览商品 → 选择规格和尺寸 → 填写地址与联系人 → 提交订单 → 查看订单进度。')
    mini_steps = {
        '首页': ['进入小程序首页，浏览 Banner、热门定制推荐和产品分类。', '点击 Banner；若 Banner 绑定了商品，会直接打开该商品详情。', '点击底部“产品”或分类入口进入产品列表。'],
        '产品列表页': ['点击分类筛选产品，或在搜索框输入产品名称/关键词。', '点击商品卡片进入详情。切换分类时，搜索框会自动清空。'],
        '产品下单页 1': ['填写门窗宽度和高度，系统会自动计算计费面积和预计价格。', '根据实际需求选择玻璃、门锁、铝材、颜色等选配。'],
        '产品下单页 2': ['点击新增方案可配置多套门窗。', '给每套方案填写套名、备注和尺寸，确认后进入提交订单页。'],
        '提交订单页': ['确认姓名、手机号和默认收货/安装地址。', '需要时上传现场图片并填写备注。', '核对方案数量、选配、面积和金额后点击提交订单。'],
        '订单列表页': ['点击状态 Tab 查看全部、待复核、生产中、待提货、已完成或已取消订单。', '状态更新提醒区域会显示最新状态 Tag 和订单号，点击订单号行或“查看详情”进入订单详情。', '点击订单卡片中的“查看详情”查看完整信息。'],
        '订单详情页': ['查看顶部当前状态和订单进度时间轴。', '向下查看客户信息、安装地址、每套定制规格、选配和金额。', '点击底部电话按钮联系门店，点击“返回订单列表”回到列表。'],
        '个人中心页': ['查看个人资料和订单状态数量。', '点击编辑个人信息、地址管理、订单状态或联系我们进入对应页面。', '点击官方电话可直接拨号。'],
        '编辑资料/登录页': ['填写姓名和手机号，必要时点击头像设置/更换头像。', '点击“保存并更新资料”。历史订单不会被修改，新资料用于后续订单。'],
        '地址列表': ['点击“新增收货/安装地址”填写地址。', '可编辑、删除地址，或点击“设为默认”切换默认地址。'],
        '联系我们': ['查看门店地址和营业时间。', '点击接单顾问电话直接拨号，按页面提示长按二维码添加顾问。'],
    }
    for i, (title, key, desc) in enumerate(MINI):
        add_section(doc, title, key, desc, mini_steps[title], mobile=True, page_break_before=bool(i))

    admin_heading = add_heading(doc, '二、管理后台使用说明', 1)
    admin_heading.paragraph_format.page_break_before = True
    add_para(doc, '管理后台用于日常运营。建议按“基础资料 → 商品配置 → 订单处理 → 权限与设置”的顺序维护，所有涉及保存、上下架和状态修改的操作都应确认页面提示。')
    admin_steps = {
        '登录页': ['打开管理后台地址，输入管理员账号和密码。', '点击“登录系统”进入后台。'],
        '大盘数据': ['查看顶部订单总数和各状态数量。', '通过趋势图、热门产品占比和订单流转图了解经营情况。', '需要最新数据时点击右上角“刷新图表数据”。'],
        'Banner 管理': ['点击“新增 Banner”，填写名称和可选副标题。', '上传 Banner 图片；需要点击后跳商品详情时，在“绑定商品”中选择商品。', '设置排序数字和启用状态，点击“确定”保存。数字越小越靠前。'],
        '门窗分类': ['点击“新建门窗分类”新增分类。', '填写分类名称、首页小标签、图标和排序权重。', '通过编辑、启用/停用和删除维护分类。'],
        '商品列表': ['使用分类、商品名称/描述筛选商品。', '点击“编辑”修改基础资料；点击“选配规则”维护加价项。', '通过“已上架/已下架”和“热门/普通”开关独立控制展示状态。'],
        '新增/编辑商品': ['填写商品名称、分类、描述、默认宽高、起步计费面积和基础单价。', '上传商品封面主图，确认上下架状态。', '点击“确定”保存，保存成功后再到商品列表核对。'],
        '商品选配规则': ['点击“新增选配分组”，选择分组名称。', '点击“添加组内选项”，填写选项名称、图片、加价方式和金额。', '按需设置默认勾选，完成后点击底部“保存选配规则配置”。'],
        '订单列表（可创建订单）': ['使用订单号、客户、商品、状态和订单时间筛选订单。', '点击“创建订单”可代客户填写方案并提交。', '点击“查看详情”查看订单；点击“修改状态”更新状态、追加费用或填写商家备注。'],
        '查看订单详情': ['查看订单状态、客户信息、收货/安装地址和方案明细。', '核对每套尺寸、面积、选配和金额。', '订单备注应以订单详情页显示的“商家订单备注”和客户方案备注为准。'],
        '接单员配置': ['点击“添加接单员”，填写姓名和联系电话。', '上传个人微信二维码，设置是否展示。', '小程序“联系我们”页面只展示启用的接单员。'],
        '客户管理（可新增/编辑/添加地址）': ['按客户昵称/姓名或联系电话搜索客户。', '点击“创建客户”新增客户，点击“修改资料”编辑客户信息。', '点击“查看详情”进入客户档案，可继续添加地址。'],
        '查看客户详情': ['查看客户 ID、微信 OpenID、身份和累计订单数。', '在“已保存的收货与安装地址”区域点击“新增地址”添加地址。', '保存后，该地址可用于后台代客户创建订单。'],
        '管理员管理': ['点击“新建管理员账号”创建后台账号。', '设置姓名/备注、联系电话、角色和账号状态。', 'ROOT 超级管理员账号受保护，不能随意删除或降权。'],
        '角色/权限': ['点击“新建角色”创建角色。', '点击“分配菜单权限”，勾选该角色允许访问的菜单。', '把角色分配给管理员后，管理员重新登录即可获得对应权限。'],
        '操作日志': ['按事件类型、管理员账号、用户昵称、手机号和日期范围筛选。', '点击“查询”查看结果，点击“重置”清空筛选条件。', '日志可用于核对订单状态、商品和系统配置是谁在何时修改的。'],
        '全局设置': ['维护门店名称、官方客服电话、门店详细地址和营业时间。', '填写腾讯地图纬度和经度，用于小程序定位。', '点击“保存全局配置”，保存成功后小程序相关页面会使用新信息。'],
    }
    for i, (title, key, desc) in enumerate(ADMIN):
        add_section(doc, title, key, desc, admin_steps[title], mobile=False, page_break_before=bool(i))

    faq_heading = add_heading(doc, '三、常见问题与注意事项', 1)
    faq_heading.paragraph_format.page_break_before = True
    add_heading(doc, '1. 小程序登录', 2)
    add_para(doc, '小程序会优先进行静默登录。同一个微信用户应保持同一账号；如果需要填写姓名、手机号或地址，按页面提示补充即可。')
    add_heading(doc, '2. 订单状态', 2)
    add_para(doc, '订单状态由管理后台更新，小程序用户不能自行取消订单。进入首页、产品页、个人中心或订单列表时，系统会检查是否有新的状态更新，并在订单 Tab 显示原生红点。')
    add_heading(doc, '3. 商品不显示', 2)
    add_para(doc, '请在管理后台确认商品已上架、所属分类已启用；首页热门区域还需要打开“热门”开关。Banner 需要启用后才会在小程序首页显示。')
    add_heading(doc, '4. 权限不足', 2)
    add_para(doc, '普通管理员只能访问所属角色授权的菜单。若提示没有权限，请让 ROOT 超级管理员在“角色/权限”中检查菜单授权，并让该管理员重新登录。')
    add_heading(doc, '5. 保存失败或加载很慢', 2)
    add_para(doc, '先确认网络正常、必填项已填写，再点击一次保存并等待结果。不要连续重复点击；若仍失败，记录页面提示和操作时间，交给系统维护人员排查。')

    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    build()
